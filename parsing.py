
import os
import re
import hashlib
from zipfile import ZipFile
from typing import List

from docx import Document
from docx.table import Table as _Tbl
from docx.text.paragraph import Paragraph as _Par

from pypdf import PdfReader

from config import (
    MAX_DOCX_ENTRIES,
    MAX_DOCX_UNCOMPRESSED_BYTES,
    MAX_PDF_PAGES,
    MAX_UPLOAD_BYTES,
    UPLOAD_TMP_DIR,
    READ_BLOCK,
)

def read_docx(file_path: str) -> str:
    with ZipFile(file_path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_DOCX_ENTRIES:
            raise ValueError("DOCX 内部文件数量超过安全上限")
        total_size = sum(entry.file_size for entry in entries)
        if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise ValueError("DOCX 解压后体积超过安全上限")
    doc = Document(file_path)

    def table_lines(tbl, prev_paras=None):

        grid = []
        for row in tbl.rows:
            cells, prev = [], None
            for c in row.cells:
                t = c.text.strip().replace("\n", " ")
                if t != prev:
                    cells.append(t)
                prev = t
            if any(cells):
                grid.append(cells)
        if not grid:
            return []

        def is_value(s):
            s = s.strip()
            if s in ("√", "×", "○", "●", "-", "~", "±", "是", "否"):
                return True
            return bool(re.match(
                r"^[<>≥≤±]?\d+(\.\d+)?([~\-]\d+(\.\d+)?)?"
                r"(%|mmol/L|mg/dL|mg|g|kg|岁|万|次|年|月|周|d|h)?$",
                s.replace(" ", "")))


        first = [x for x in grid[0] if x]
        second = [x for x in grid[1] if x] if len(grid) >= 2 else []
        caption, header, body = "", [], grid
        if (len(first) <= 2 and len(second) > len(first)
                and not any(is_value(x) for x in second)):

            caption = " ".join(first)
            header, body = grid[1], grid[2:]
        elif len(grid) >= 2 and not any(is_value(x) for x in first):

            header, body = grid[0], grid[1:]


        if header and body:
            nxt = [x for x in body[0] if x]
            if len(nxt) > len(header) and not any(is_value(x) for x in nxt):
                header, body = body[0], body[1:]


        if not caption and prev_paras:
            cand = ""
            for p in reversed(prev_paras[-3:]):
                if re.match(r"^(表|附录|续)\s*\d*", p):
                    cand = p
                    break
            if not cand:
                cand = prev_paras[-1]
            if 4 <= len(cand) <= 60 and not cand.endswith(("。", ".", "！", "？")):
                cjk = sum(1 for ch in cand if "\u4e00" <= ch <= "\u9fff")

                if cjk >= 4 or re.match(r"^(Table|Figure|Appendix|Tab\.?)\s*\d+", cand, re.I):
                    caption = cand

        if len(caption) > 40:
            caption = caption[:40] + "…"
        tag = ("【表\u00b7%s】" % caption) if caption else "【表】"

        out = []
        for r in body:
            if len(header) >= 2 and 0 < len(r) <= len(header):
                out.append(tag + "；".join(
                    ("%s：%s" % (header[i], r[i])) if header[i] else r[i]
                    for i in range(len(r))))
            else:
                out.append(tag + " | ".join(r))
        return out

    parts = []
    try:

        from docx.table import Table as _Tbl
        from docx.text.paragraph import Paragraph as _Par

        prev_paras = []
        for child in doc.element.body.iterchildren():
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                t = _Par(child, doc).text.strip()
                if t:
                    parts.append(t)
                    prev_paras.append(t)
                    del prev_paras[:-4]
            elif tag == 'tbl':
                tl = table_lines(_Tbl(child, doc), prev_paras)
                if tl:
                    parts.append("[表格]\n" + "\n".join(tl) + "\n[/表格]")
    except Exception as e:

        print("【警告】docx 顺序遍历失败，退回「先段落、后表格」模式：%s: %s"
              % (type(e).__name__, e))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for tbl in doc.tables:
            tl = table_lines(tbl)
            if tl:
                parts.append("[表格]\n" + "\n".join(tl) + "\n[/表格]")

    return "\n".join(parts)

def read_txt(file_path: str) -> str:
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'utf-16'):
        try:
            with open(file_path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()

def read_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError("PDF 页数超过安全上限")
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""
    return text

def chinese_rag_slice_advanced(
    text: str,
    chunk_size: int = 400,
    overlap_ratio: float = 0.15,
    min_slice_len: int = 20
) -> List[str]:
    if not text or not text.strip():
        return []


    text = text.replace('[表格]', '').replace('[/表格]', '')


    table_rows = []
    non_table_lines = []
    for line in text.split('\n'):
        line = line.strip()
        if not line:
            continue


        if line.startswith('【表'):

            if len(line) >= 5:
                table_rows.append(line)
        else:
            non_table_lines.append(line)


    body_text = ' '.join(non_table_lines)
    body_text = re.sub(r'\s+', ' ', body_text).strip()

    final_slices = []

    if body_text:

        zh_end = r'。！？；…'
        sentences = re.split(rf'(?<=[{zh_end}])|(?<=[.!?;])(?=\s|$)', body_text)
        sentences = [s.strip() for s in sentences if s.strip()]
        
        if not sentences:
            return table_rows
        

        overlap_size = int(chunk_size * overlap_ratio)
        current_chunk = []
        current_length = 0
        
        for sentence in sentences:
            sentence_len = len(sentence)
            

            if sentence_len >= chunk_size:
                if current_chunk:
                    final_slices.append(' '.join(current_chunk))
                    current_chunk = []
                    current_length = 0
                final_slices.append(sentence)
                continue
            

            if current_length + sentence_len > chunk_size:
                final_slices.append(' '.join(current_chunk))
                

                overlap_sentences = []
                overlap_length = 0
                for s in reversed(current_chunk):
                    if overlap_length + len(s) > overlap_size:
                        break
                    overlap_sentences.insert(0, s)
                    overlap_length += len(s)
                
                current_chunk = overlap_sentences
                current_length = overlap_length
            
            current_chunk.append(sentence)
            current_length += sentence_len
        

        if current_chunk:
            final_slices.append(' '.join(current_chunk))
    

    for row in table_rows:
        if row not in final_slices:
            final_slices.append(row)

    return final_slices


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.docx':
        return read_docx(file_path)
    if ext == '.pdf':
        return read_pdf(file_path)
    return read_txt(file_path)

def text_hash_of(text: str) -> str:
    normalized = re.sub(r'\s+', '', text)
    return hashlib.md5(normalized.encode('utf-8')).hexdigest()[:16]
